"""Extract bounded worksheet cells with exact sheet, row and column locators."""

from io import BytesIO
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

from app.domain.sources import Locator
from app.infrastructure.ingestion.csv_table import (
    MAX_CSV_CELLS,
    MAX_CSV_COLUMNS,
    MAX_CSV_FIELD_CHARS,
    MAX_CSV_LABEL_CHARS,
    MAX_CSV_ROWS,
    ExtractedCell,
)
from app.infrastructure.ingestion.pdf_text import ExtractionError
from app.infrastructure.ingestion.storage import MAX_XLSX_BYTES

MAX_XLSX_SHEETS = 8
MAX_ZIP_PARTS = 200
MAX_UNCOMPRESSED_BYTES = 20 * 1024 * 1024


def _cell_text(value: object) -> str:
    return "" if value is None else str(value)


def extract_xlsx_cells(data: bytes) -> tuple[ExtractedCell, ...]:
    if not data or len(data) > MAX_XLSX_BYTES:
        raise ExtractionError("invalid XLSX size")
    try:
        with ZipFile(BytesIO(data)) as archive:
            parts = archive.infolist()
            if not 1 <= len(parts) <= MAX_ZIP_PARTS or sum(
                part.file_size for part in parts
            ) > MAX_UNCOMPRESSED_BYTES:
                raise ExtractionError("XLSX archive exceeds limits")
            if any(
                part.flag_bits & 1
                or part.file_size > MAX_UNCOMPRESSED_BYTES
                or "externalLinks/" in part.filename
                or part.filename.endswith("vbaProject.bin")
                or part.filename.startswith((
                    "xl/drawings/", "xl/media/", "xl/charts/", "xl/pivot",
                ))
                for part in parts
            ) or len({part.filename for part in parts}) != len(parts):
                raise ExtractionError("XLSX archive contains unsupported parts")
        workbook = load_workbook(BytesIO(data), read_only=False, data_only=False,
                                 keep_links=False)
        try:
            if not 1 <= len(workbook.worksheets) <= MAX_XLSX_SHEETS:
                raise ExtractionError("XLSX sheet count is unsupported")
            if len(workbook.sheetnames) != len(workbook.worksheets):
                raise ExtractionError("XLSX chart sheets are unsupported")
            cells: list[ExtractedCell] = []
            for sheet in workbook.worksheets:
                if sheet.sheet_state != "visible":
                    raise ExtractionError("hidden XLSX sheets need review")
                if not 2 <= sheet.max_column <= MAX_CSV_COLUMNS or sheet.max_row > MAX_CSV_ROWS:
                    raise ExtractionError("XLSX dimensions exceed limits")
                if any(dimension.hidden for dimension in sheet.row_dimensions.values()) or any(
                    dimension.hidden for dimension in sheet.column_dimensions.values()
                ):
                    raise ExtractionError("hidden XLSX rows or columns need review")
                merged_keys: dict[int, str] = {}
                for span in sheet.merged_cells.ranges:
                    if (span.min_col != 1 or span.max_col != 1 or span.min_row < 2
                            or span.max_row > MAX_CSV_ROWS):
                        raise ExtractionError("XLSX merged range has ambiguous table meaning")
                    key = _cell_text(sheet.cell(span.min_row, 1).value).strip()
                    if not key or len(key) > MAX_CSV_LABEL_CHARS:
                        raise ExtractionError("XLSX merged row key is invalid")
                    for row_number in range(span.min_row, span.max_row + 1):
                        merged_keys[row_number] = key
                header = [_cell_text(sheet.cell(1, column).value).strip()
                          for column in range(1, sheet.max_column + 1)]
                if (any(not item or len(item) > MAX_CSV_LABEL_CHARS for item in header)
                        or len(set(header)) != len(header)):
                    raise ExtractionError("XLSX needs unique, nonempty headers")
                for row_number in range(1, sheet.max_row + 1):
                    for column in range(1, sheet.max_column + 1):
                        cell = sheet.cell(row_number, column)
                        if cell.data_type in {"f", "e"}:
                            raise ExtractionError("XLSX formulas and errors need review")
                for row_number in range(2, sheet.max_row + 1):
                    key = merged_keys.get(
                        row_number, _cell_text(sheet.cell(row_number, 1).value).strip()
                    )
                    row_values = [
                        _cell_text(sheet.cell(row_number, column).value).strip()
                        for column in range(2, sheet.max_column + 1)
                    ]
                    if not key and not any(row_values):
                        continue
                    if not key or len(key) > MAX_CSV_LABEL_CHARS:
                        raise ExtractionError("XLSX row key is invalid")
                    for column, value in enumerate(row_values, start=2):
                        if len(value) > MAX_CSV_FIELD_CHARS:
                            raise ExtractionError("XLSX cell exceeds limit")
                        if not value:
                            continue
                        cells.append(ExtractedCell(
                            f"{header[0]}: {key}; {header[column - 1]}: {value}",
                            Locator(sheet=sheet.title, row_start=row_number,
                                    row_end=row_number, column_start=column,
                                    column_end=column),
                        ))
                        if len(cells) > MAX_CSV_CELLS:
                            raise ExtractionError("XLSX exceeds cell limit")
            if not cells:
                raise ExtractionError("XLSX has no queryable data cells")
            return tuple(cells)
        finally:
            workbook.close()
    except ExtractionError:
        raise
    except (BadZipFile, OSError, ValueError, TypeError, KeyError) as exc:
        raise ExtractionError("XLSX decoding failed") from exc
