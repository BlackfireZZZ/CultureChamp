"""Self-authored workbook used only to verify table mechanics."""

from io import BytesIO

from openpyxl import Workbook


def self_authored_xlsx() -> bytes:
    workbook = Workbook()
    first = workbook.active
    first.title = "North"
    first.append(["Item", "Count", "Note"])
    first.append(["Example A", None, "Seven bells"])
    first.append([None, 7, None])
    first.merge_cells("A2:A3")
    second = workbook.create_sheet("South")
    second.append(["Item", "Count"])
    second.append(["Example B", 4])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def self_authored_formula_xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Item", "Count"])
    sheet.append(["Example", "=2+2"])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
